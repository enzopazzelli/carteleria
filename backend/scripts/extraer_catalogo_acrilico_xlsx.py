"""Extrae del export real de AppSheet (`CARTELERIA 2026.xlsx`) el
catálogo de la categoría `ACRILICOS` — acrílico, policarbonato, alto
impacto y corrugado plástico, tal como el propio inventario del cliente
los agrupa. Hermano de `extraer_catalogo_chapa_xlsx.py`, mismo criterio
de preparación para desarrollo local (no reemplaza `CART-104`).

**El xlsx de entrada y el JSON de salida nunca van al repositorio** —
mismo criterio que el de chapa (PII y precios reales del cliente).

Uso:
    python scripts/extraer_catalogo_acrilico_xlsx.py --xlsx "../CARTELERIA 2026.xlsx" --out local/catalogo_acrilico.json

Por qué un script aparte y no reutilizar el de chapa: la descripción de
esta categoría no trae `cal. N` (calibre entero) sino el espesor en mm
directo, a veces con coma decimal y espacio irregular antes de "mm"
(`"... - 2mm"` vs. `"... - 2,4 mm"`) — un regex distinto, no el mismo
con un parámetro.
"""
from __future__ import annotations

import argparse
import json
import re
from decimal import Decimal
from pathlib import Path

import openpyxl

from _precios_cotizador import leer_precios_cotizador

_PATRON_FORMATO = re.compile(
    r"(?P<material>.+?)\s*-\s*(?P<ancho>\d+[.,]\d+)\s*x\s*(?P<alto>\d+[.,]\d+)\s*-\s*"
    r"(?P<espesor>\d+(?:[.,]\d+)?)\s*mm",
    re.IGNORECASE,
)
_M_A_MM = Decimal(1000)


def _decimal_es(texto: str) -> Decimal:
    return Decimal(texto.replace(",", "."))


def _leer_formatos_acrilico(ws) -> tuple[list[dict], list[dict]]:
    encabezado = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    formatos, sin_parsear = [], []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        registro = dict(zip(encabezado, fila))
        if str(registro.get("CATEGORIA", "")).strip().upper() != "ACRILICOS":
            continue
        descripcion = registro.get("DESCRIPCION") or ""
        match = _PATRON_FORMATO.search(descripcion)
        if not match:
            sin_parsear.append({"codigo": registro.get("CODIGO"), "descripcion": descripcion})
            continue
        formatos.append(
            {
                "codigo": registro.get("CODIGO"),
                "material": match.group("material").strip(),
                "espesor_mm": str(_decimal_es(match.group("espesor"))),
                "ancho_mm": str(_decimal_es(match.group("ancho")) * _M_A_MM),
                "alto_mm": str(_decimal_es(match.group("alto")) * _M_A_MM),
                "costo_unidad_venta": None,
                "unidad_venta": None,
                "moneda": None,
            }
        )
    return formatos, sin_parsear


def _completar_costos(formatos: list[dict], ws_cotizador) -> None:
    precios = leer_precios_cotizador(ws_cotizador)
    for formato in formatos:
        precio = precios.get(formato["codigo"])
        if precio is None:
            continue
        formato["costo_unidad_venta"] = precio["costo_unidad_venta"]
        formato["unidad_venta"] = precio["unidad_venta"]
        formato["moneda"] = precio["moneda"]


def extraer_catalogo(ruta_xlsx: Path) -> dict:
    libro = openpyxl.load_workbook(ruta_xlsx, read_only=True, data_only=True)
    formatos, sin_parsear = _leer_formatos_acrilico(libro["INVENTARIO"])
    _completar_costos(formatos, libro["COTIZADOR"])

    sin_precio = [f["codigo"] for f in formatos if f["costo_unidad_venta"] is None]
    advertencias = [
        f"{len(sin_parsear)} formato(s) de acrílico no se pudieron parsear desde DESCRIPCION.",
        f"{len(sin_precio)} de {len(formatos)} formato(s) no tienen costo real en COTIZADOR "
        "(ausentes en la planilla, o con costo en 0 sin margen configurado): "
        + ", ".join(sin_precio),
    ]
    return {"formatos": formatos, "sin_parsear": sin_parsear, "advertencias": advertencias}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--xlsx", required=True, type=Path, help="Ruta al xlsx real del cliente")
    parser.add_argument("--out", type=Path, default=None, help="Archivo de salida (default: stdout)")
    args = parser.parse_args()

    catalogo = extraer_catalogo(args.xlsx)
    salida = json.dumps(catalogo, indent=2, ensure_ascii=False)

    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(salida, encoding="utf-8")
        print(f"Catálogo escrito en {args.out} — no lo commitees (ver docstring del script).")
    else:
        print(salida)


if __name__ == "__main__":
    main()
